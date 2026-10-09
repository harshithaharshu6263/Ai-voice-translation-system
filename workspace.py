"""Authenticated translation workspace with per-user SQLite storage."""
import base64
import functools
import io
import os
from pathlib import Path
import re
import secrets
import sqlite3
import subprocess
import tempfile
import time
import wave
import requests

from flask import Flask, g, jsonify, redirect, render_template, request, send_file, session, url_for
from flask_mail import Mail, Message
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.exceptions import HTTPException
from deep_translator import GoogleTranslator
from dotenv import load_dotenv
from langdetect import detect, DetectorFactory, LangDetectException

BASE = Path(__file__).resolve().parent
load_dotenv(BASE / '.env')
DetectorFactory.seed = 0
EMAIL_RE = re.compile(r'^[^\s@]+@[^\s@]+\.[^\s@]+$')


def translate_online(text, source, target):
    """Google's web endpoint, with bounded requests and provider detection."""
    if source == target:
        return text, source
    response = requests.get('https://translate.googleapis.com/translate_a/single',
                            params={'client': 'gtx', 'sl': source, 'tl': target, 'dt': 't', 'q': text},
                            timeout=(5, 20))
    response.raise_for_status()
    result = response.json()
    translated = ''.join(part[0] for part in result[0] if part and isinstance(part[0], str))
    detected = result[2] if len(result) > 2 and isinstance(result[2], str) else None
    return translated, detected


def create_app(test_config=None):
    app = Flask(__name__, instance_path=str(BASE / 'instance'))
    Path(app.instance_path).mkdir(exist_ok=True)
    secret_file = Path(app.instance_path) / 'secret.key'
    if not secret_file.exists():
        secret_file.write_text(secrets.token_hex(32), encoding='utf-8')
    app.config.update(
        SECRET_KEY=os.getenv('SECRET_KEY') or secret_file.read_text().strip(),
        DATABASE=str(Path(app.instance_path) / 'voice_translate.sqlite3'),
        MAX_CONTENT_LENGTH=25 * 1024 * 1024,
        SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE='Lax',
        SESSION_COOKIE_SECURE=os.getenv('COOKIE_SECURE', 'false').lower() == 'true',
        MAIL_SERVER=os.getenv('MAIL_SERVER', 'smtp.gmail.com'),
        MAIL_PORT=int(os.getenv('MAIL_PORT', '587')),
        MAIL_USE_TLS=os.getenv('MAIL_USE_TLS', 'true').lower() == 'true',
        MAIL_USE_SSL=os.getenv('MAIL_USE_SSL', 'false').lower() == 'true',
        MAIL_USERNAME=os.getenv('MAIL_USERNAME', ''), MAIL_PASSWORD=os.getenv('MAIL_PASSWORD', ''),
        MAIL_DEFAULT_SENDER=os.getenv('MAIL_DEFAULT_SENDER') or os.getenv('MAIL_USERNAME', ''),
        MAIL_TIMEOUT=20,
    )
    if test_config:
        app.config.update(test_config)
    mail = Mail(app)
    attempts = {}
    languages = GoogleTranslator().get_supported_languages(as_dict=True)
    codes = set(languages.values())

    def db():
        if 'db' not in g:
            g.db = sqlite3.connect(app.config['DATABASE'])
            g.db.row_factory = sqlite3.Row
            g.db.execute('PRAGMA foreign_keys = ON')
        return g.db

    @app.teardown_appcontext
    def close_db(error=None):
        connection = g.pop('db', None)
        if connection:
            connection.close()

    with app.app_context():
        db().executescript('''
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY, email TEXT UNIQUE NOT NULL,
                name TEXT NOT NULL, password_hash TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS translations (
                id INTEGER PRIMARY KEY, user_id INTEGER NOT NULL REFERENCES users(id),
                original TEXT NOT NULL, translated TEXT NOT NULL,
                source TEXT NOT NULL, target TEXT NOT NULL, detected TEXT,
                favorite INTEGER NOT NULL DEFAULT 0, mode TEXT NOT NULL DEFAULT 'translate',
                speaker TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL DEFAULT (strftime('%Y-%m-%dT%H:%M:%fZ','now'))
            );
            CREATE INDEX IF NOT EXISTS history_user ON translations(user_id, id DESC);
        ''')
        if os.getenv('DEMO_LOGIN', 'true').lower() == 'true' and not test_config:
            db().execute('INSERT OR IGNORE INTO users(email,name,password_hash) VALUES (?,?,?)',
                         ('harshita@gmail.com', 'Harshita', generate_password_hash('1234')))
        db().commit()

    def csrf_token():
        if 'csrf' not in session:
            session['csrf'] = secrets.token_hex(32)
        return session['csrf']

    app.jinja_env.globals['csrf_token'] = csrf_token

    @app.before_request
    def protect_requests():
        if request.method in ('POST', 'PUT', 'PATCH', 'DELETE'):
            token = request.headers.get('X-CSRF-Token') or request.form.get('csrf_token', '')
            if not session.get('csrf') or not secrets.compare_digest(token, session['csrf']):
                return jsonify(error='Your session expired. Refresh the page and try again.'), 403

    @app.after_request
    def headers(response):
        response.headers['X-Content-Type-Options'] = 'nosniff'
        response.headers['X-Frame-Options'] = 'SAMEORIGIN'
        if not request.path.startswith('/static/'):
            response.headers['Cache-Control'] = 'no-store'
        return response

    @app.errorhandler(HTTPException)
    def http_error(error):
        if error.code == 413:
            return jsonify(error='The upload is too large. Choose a file under 25 MB.'), 413
        return jsonify(error=error.description), error.code

    def auth(fn):
        @functools.wraps(fn)
        def wrapped(*args, **kwargs):
            if not session.get('user_id'):
                return jsonify(error='Please sign in again.'), 401
            return fn(*args, **kwargs)
        return wrapped

    def payload():
        data = request.get_json(silent=True)
        return data if isinstance(data, dict) else {}

    def text_value(data, key, maximum=5000):
        value = data.get(key, '')
        if not isinstance(value, str) or not value.strip() or len(value) > maximum:
            raise ValueError(f'{key.replace("_", " ").capitalize()} must contain 1–{maximum} characters.')
        return value.strip()

    def owned_record(record_id):
        return db().execute('SELECT * FROM translations WHERE id=? AND user_id=?',
                            (record_id, session['user_id'])).fetchone()

    @app.route('/', methods=['GET', 'POST'])
    @app.route('/register', methods=['GET', 'POST'])
    def login():
        register = request.path == '/register'
        error = None
        if session.get('user_id'):
            return redirect(url_for('dashboard'))
        if request.method == 'POST':
            email = request.form.get('email', '').strip().lower()
            password = request.form.get('password', '')
            name = request.form.get('name', '').strip()
            key = (request.remote_addr, email)
            failures = [t for t in attempts.get(key, []) if time.monotonic() - t < 900]
            if len(failures) >= 8:
                error = 'Too many attempts. Try again in 15 minutes.'
            elif not EMAIL_RE.fullmatch(email) or len(email) > 254:
                error = 'Enter a valid email address.'
            elif register:
                if not name or len(name) > 80:
                    error = 'Enter a name of up to 80 characters.'
                elif not 8 <= len(password) <= 128:
                    error = 'Choose a password with 8–128 characters.'
                else:
                    try:
                        cursor = db().execute('INSERT INTO users(email,name,password_hash) VALUES (?,?,?)',
                                              (email, name, generate_password_hash(password)))
                        db().commit()
                        session.clear()
                        session['user_id'] = cursor.lastrowid
                        return redirect(url_for('dashboard'))
                    except sqlite3.IntegrityError:
                        error = 'This email is already registered. Sign in instead.'
            else:
                user = db().execute('SELECT * FROM users WHERE email=?', (email,)).fetchone()
                if user and len(password) <= 128 and check_password_hash(user['password_hash'], password):
                    session.clear()
                    session['user_id'] = user['id']
                    attempts.pop(key, None)
                    return redirect(url_for('dashboard'))
                error = 'Email or password is incorrect.'
            if error:
                failures.append(time.monotonic())
                attempts[key] = failures
        return render_template('auth.html', register=register, error=error)

    @app.get('/dashboard')
    def dashboard():
        if not session.get('user_id'):
            return redirect(url_for('login'))
        user = db().execute('SELECT id,email,name FROM users WHERE id=?', (session['user_id'],)).fetchone()
        if not user:
            session.clear()
            return redirect(url_for('login'))
        return render_template('workspace.html', user=user)

    @app.post('/logout')
    def logout():
        session.clear()
        return redirect(url_for('login'))

    @app.get('/api/config')
    @auth
    def configuration():
        from gtts.lang import tts_langs
        return jsonify(languages=languages, tts_languages=tts_langs(),
                       email_configured=bool(app.config['MAIL_USERNAME'] and app.config['MAIL_PASSWORD']))

    @app.route('/api/profile', methods=['GET', 'PATCH'])
    @auth
    def profile():
        if request.method == 'PATCH':
            data = payload()
            try:
                name = text_value(data, 'name', 80)
                if data.get('new_password'):
                    user = db().execute('SELECT * FROM users WHERE id=?', (session['user_id'],)).fetchone()
                    current, new = data.get('current_password', ''), data['new_password']
                    if not isinstance(current, str) or not check_password_hash(user['password_hash'], current):
                        return jsonify(error='Current password is incorrect.'), 400
                    if not isinstance(new, str) or not 8 <= len(new) <= 128:
                        return jsonify(error='Choose a password with 8–128 characters.'), 400
                    db().execute('UPDATE users SET password_hash=? WHERE id=?',
                                 (generate_password_hash(new), session['user_id']))
                db().execute('UPDATE users SET name=? WHERE id=?', (name, session['user_id']))
                db().commit()
            except ValueError as error:
                return jsonify(error=str(error)), 400
        user = db().execute('SELECT id,email,name FROM users WHERE id=?', (session['user_id'],)).fetchone()
        return jsonify(dict(user))

    @app.post('/translate')
    @auth
    def translate_text():
        data = payload()
        try:
            text = text_value(data, 'text')
            source, target = data.get('source', 'auto'), data.get('target', 'en')
            if not isinstance(target, str) or not isinstance(source, str) or target not in codes or (source != 'auto' and source not in codes):
                return jsonify(error='Choose a supported language.'), 400
            detected = None
            try:
                detected = detect(text)
                detected = {'zh-cn': 'zh-CN', 'zh-tw': 'zh-TW', 'he': 'iw'}.get(detected, detected)
            except LangDetectException:
                pass
            translated, provider_detected = translate_online(text, source, target)
            if provider_detected:
                detected = {'he': 'iw', 'zh-cn': 'zh-CN', 'zh-tw': 'zh-TW'}.get(provider_detected, provider_detected)
            if not translated:
                return jsonify(error='No translation returned. Try again.'), 502
            mode = 'conversation' if data.get('mode') == 'conversation' else 'translate'
            speaker = data.get('speaker', '') if mode == 'conversation' else ''
            if speaker not in ('', 'A', 'B'):
                return jsonify(error='Invalid speaker.'), 400
            cursor = db().execute('''INSERT INTO translations
                (user_id, original, translated, source, target, detected, mode, speaker)
                VALUES (?,?,?,?,?,?,?,?)''',
                (session['user_id'], text, translated, source, target, detected, mode, speaker))
            db().commit()
            return jsonify(dict(owned_record(cursor.lastrowid)))
        except ValueError as error:
            return jsonify(error=str(error)), 400
        except requests.Timeout:
            return jsonify(error='Translation timed out. Check your connection and try again.'), 504
        except requests.HTTPError as error:
            if error.response is not None and error.response.status_code == 429:
                return jsonify(error='The translation service is busy. Wait a moment and try again.'), 429
            return jsonify(error='The translation service is unavailable. Try again later.'), 502
        except Exception:
            app.logger.exception('Translation service failed')
            return jsonify(error='The translation service is unavailable. Check your connection and try again.'), 502

    @app.get('/api/history')
    @auth
    def history():
        query = request.args.get('q', '').strip()[:200]
        favorite = request.args.get('favorites') == '1'
        try:
            before = max(0, int(request.args.get('before', '0')))
        except ValueError:
            return jsonify(error='Invalid history cursor.'), 400
        sql = 'SELECT * FROM translations WHERE user_id=?'
        params = [session['user_id']]
        if query:
            sql += ' AND (instr(lower(original),lower(?))>0 OR instr(lower(translated),lower(?))>0)'
            params += [query, query]
        if favorite:
            sql += ' AND favorite=1'
        if before:
            sql += ' AND id<?'
            params.append(before)
        rows = db().execute(sql + ' ORDER BY id DESC LIMIT 51', params).fetchall()
        return jsonify(items=[dict(row) for row in rows[:50]], has_more=len(rows) > 50)

    @app.route('/api/history/<int:record_id>', methods=['PATCH', 'DELETE'])
    @auth
    def change_history(record_id):
        if not owned_record(record_id):
            return jsonify(error='Translation not found.'), 404
        if request.method == 'DELETE':
            db().execute('DELETE FROM translations WHERE id=? AND user_id=?', (record_id, session['user_id']))
        else:
            favorite = payload().get('favorite')
            if not isinstance(favorite, bool):
                return jsonify(error='Favorite must be true or false.'), 400
            db().execute('UPDATE translations SET favorite=? WHERE id=? AND user_id=?',
                         (int(favorite), record_id, session['user_id']))
        db().commit()
        return jsonify(success=True)

    @app.post('/api/transcribe')
    @auth
    def transcribe():
        import imageio_ffmpeg
        import speech_recognition as sr
        upload = request.files.get('audio')
        language = request.form.get('language', 'en-IN')
        if not upload or not upload.filename:
            return jsonify(error='Choose an audio file.'), 400
        if Path(upload.filename).suffix.lower() not in ('.wav', '.mp3', '.m4a', '.ogg', '.webm', '.flac', '.aac'):
            return jsonify(error='Upload WAV, MP3, M4A, OGG, WebM, FLAC, or AAC audio.'), 400
        if not re.fullmatch(r'[a-zA-Z]{2,3}(?:-[a-zA-Z]{2,4})?', language):
            return jsonify(error='Choose a valid spoken language.'), 400
        try:
            with tempfile.TemporaryDirectory() as temp:
                original = Path(temp) / ('input' + Path(upload.filename).suffix.lower())
                wav = Path(temp) / 'audio.wav'
                upload.save(original)
                result = subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), '-nostdin', '-v', 'error',
                    '-y', '-i', str(original), '-t', '121', '-vn', '-ac', '1', '-ar', '16000', str(wav)],
                    capture_output=True, timeout=40,
                    creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
                if result.returncode != 0:
                    return jsonify(error='This file could not be decoded. Choose a valid audio recording.'), 400
                with wave.open(str(wav), 'rb') as audio:
                    duration = audio.getnframes() / audio.getframerate()
                if duration > 120:
                    return jsonify(error='Choose a recording of up to 2 minutes.'), 400
                recognizer = sr.Recognizer()
                recognizer.operation_timeout = 25
                parts = []
                with sr.AudioFile(str(wav)) as audio:
                    for _ in range(max(1, int(duration // 25) + 1)):
                        chunk = recognizer.record(audio, duration=25)
                        if not chunk.frame_data:
                            break
                        try:
                            parts.append(recognizer.recognize_google(chunk, language=language,
                                endpoint='https://www.google.com/speech-api/v2/recognize'))
                        except sr.UnknownValueError:
                            continue
                if not parts:
                    return jsonify(error='No speech was recognized. Try a clearer recording and check the spoken language.'), 422
                return jsonify(text=' '.join(parts), duration=round(duration, 1))
        except subprocess.TimeoutExpired:
            return jsonify(error='Audio conversion timed out. Try a smaller recording.'), 504
        except sr.RequestError:
            return jsonify(error='Speech recognition is unavailable. Check your connection and try again.'), 502
        except Exception:
            app.logger.exception('Audio transcription failed')
            return jsonify(error='Unable to transcribe this recording. Try another audio file.'), 502

    def make_audio(text, language):
        from gtts import gTTS
        from gtts.lang import tts_langs
        if not isinstance(language, str) or language not in tts_langs():
            raise ValueError('Downloadable speech is unavailable in this language. Try browser playback.')
        buffer = io.BytesIO()
        gTTS(text=text, lang=language, timeout=(10, 25)).write_to_fp(buffer)
        return buffer.getvalue()

    @app.post('/tts-generate')
    @auth
    def tts_generate():
        data = payload()
        try:
            audio = make_audio(text_value(data, 'text'), data.get('lang', 'en'))
            return jsonify(success=True, audio_base64=base64.b64encode(audio).decode())
        except ValueError as error:
            return jsonify(error=str(error)), 400
        except Exception:
            app.logger.exception('Speech generation failed')
            return jsonify(error='Audio generation is unavailable. Try browser playback or try again later.'), 502

    def word_document(record):
        from docx import Document
        document = Document()
        document.add_heading('VoiceTranslate', 0)
        document.add_paragraph(record['created_at'])
        document.add_heading('Original', level=1)
        document.add_paragraph(record['original'])
        document.add_heading('Translation', level=1)
        document.add_paragraph(record['translated'])
        buffer = io.BytesIO()
        document.save(buffer)
        buffer.seek(0)
        return buffer

    @app.get('/api/export/<int:record_id>/<format>')
    @auth
    def export(record_id, format):
        record = owned_record(record_id)
        if not record:
            return jsonify(error='Translation not found.'), 404
        if format == 'txt':
            content = f'VoiceTranslate\n{record["source"]} → {record["target"]}\n{record["created_at"]}\n\nOriginal\n{record["original"]}\n\nTranslation\n{record["translated"]}\n'
            buffer, mimetype = io.BytesIO(content.encode('utf-8-sig')), 'text/plain; charset=utf-8'
        elif format == 'docx':
            buffer, mimetype = word_document(record), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        else:
            return jsonify(error='Choose a text or Word export. Use Print / PDF for a PDF.'), 400
        return send_file(buffer, mimetype=mimetype, as_attachment=True,
                         download_name=f'translation-{record_id}.{format}')

    @app.post('/api/email')
    @auth
    def email_translation():
        if not app.config['MAIL_USERNAME'] or not app.config['MAIL_PASSWORD']:
            return jsonify(error='Email is not configured yet. Add SMTP settings to .env and restart the app.'), 503
        data = request.form if request.mimetype == 'multipart/form-data' else payload()
        recipient = data.get('email', '')
        if not isinstance(recipient, str) or len(recipient) > 254 or not EMAIL_RE.fullmatch(recipient):
            return jsonify(error='Enter a valid recipient email address.'), 400
        record_id = data.get('id')
        if request.mimetype == 'multipart/form-data':
            try:
                record_id = int(record_id)
            except (TypeError, ValueError):
                record_id = None
        if not isinstance(record_id, int):
            return jsonify(error='Translate or open a saved translation first.'), 400
        record = owned_record(record_id)
        if not record:
            return jsonify(error='Translation not found.'), 404
        kind = data.get('attachment', 'none')
        if kind not in ('none', 'audio', 'docx', 'pdf', 'recording'):
            return jsonify(error='Choose a supported attachment.'), 400
        try:
            msg = Message(subject='Your VoiceTranslate translation', recipients=[recipient])
            msg.body = f'Original:\n{record["original"]}\n\nTranslation:\n{record["translated"]}'
            if kind == 'audio':
                msg.attach('translation.mp3', 'audio/mpeg', make_audio(record['translated'], record['target']))
            elif kind == 'docx':
                msg.attach('translation.docx', 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', word_document(record).getvalue())
            elif kind in ('pdf', 'recording'):
                attachment = request.files.get('attachment_file')
                if not attachment or not attachment.filename:
                    return jsonify(error='Choose the file to attach.'), 400
                suffix = Path(attachment.filename).suffix.lower()
                content = attachment.read()
                if not content:
                    return jsonify(error='The attachment is empty.'), 400
                if kind == 'pdf':
                    if suffix != '.pdf' or not content.startswith(b'%PDF-'):
                        return jsonify(error='Choose a valid PDF. Save one using Print / PDF first.'), 400
                    msg.attach('translation.pdf', 'application/pdf', content)
                else:
                    types = {'.mp3': 'audio/mpeg', '.wav': 'audio/wav', '.m4a': 'audio/mp4',
                             '.webm': 'audio/webm', '.ogg': 'audio/ogg', '.flac': 'audio/flac', '.aac': 'audio/aac'}
                    if suffix not in types:
                        return jsonify(error='Choose a supported audio recording.'), 400
                    msg.attach('recording' + suffix, types[suffix], content)
            mail.send(msg)
            return jsonify(success=True)
        except ValueError as error:
            return jsonify(error=str(error)), 400
        except Exception:
            app.logger.exception('Email delivery failed')
            return jsonify(error='Email could not be sent. Check SMTP credentials and your connection.'), 502

    return app
